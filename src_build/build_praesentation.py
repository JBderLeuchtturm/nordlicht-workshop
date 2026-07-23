
"""Erzeugt Bilder, Titel und Sprechernotizen für docs/praesentation.html aus der PPTX.

Immer dann ausführen, wenn sich RAG_Advanced_Folien.pptx geändert hat:

    python src_build/build_praesentation.py

Titel, Kapitel und Notizen werden automatisch aus der Datei gelesen — Folien dürfen
also hinzukommen, wegfallen oder umbenannt werden, ohne dass hier etwas anzupassen ist.

Voraussetzungen: LibreOffice (soffice), poppler-utils (pdftoppm), Pillow, python-pptx.
Ohne LibreOffice: Folien in PowerPoint über "Datei > Exportieren > Bilder" als JPG mit
1600 px Breite ausgeben, als 01.jpg … NN.jpg nach docs/assets/folien/ legen und dieses
Skript mit --nur-notizen aufrufen.
"""
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "RAG_Advanced_Folien.pptx"
OUT_IMG = ROOT / "docs" / "assets" / "folien"
PRESENTER = ROOT / "docs" / "praesentation.html"
TMP = ROOT / ".build_praes"
BREITE = 1600
QUALITAET = 90


# PowerPoint ersetzt getippte Pfeile durch Wingdings-Zeichen aus der Private Use Area.
# Auf der Folie sieht das richtig aus, im HTML-Titel waere es Datenmuell.
SONDERZEICHEN = {
    "\uf0e0": "\u2192", "\uf0df": "\u2190", "\uf0e1": "\u2191", "\uf0e2": "\u2193",
    "\uf0b7": "\u00b7", "\uf0a7": "\u00b7", "\uf0fc": "\u2713", "\uf0fb": "\u2717",
    "\uf0d8": "\u25b6", "\uf0e8": "\u21d2",
}


def saeubern(text: str) -> str:
    for alt, neu in SONDERZEICHEN.items():
        text = text.replace(alt, neu)
    # was sonst noch aus der Private Use Area kommt, faellt weg
    text = "".join(c for c in text if not "\ue000" <= c <= "\uf8ff")
    return re.sub(r"\s+", " ", text).strip()


def kicker(folie) -> str:
    """Kleine Grossbuchstaben-Zeile ueber dem Titel, z. B. 'ZUM MITNEHMEN'."""
    for shape in folie.shapes:
        if not shape.has_text_frame:
            continue
        text = saeubern(shape.text_frame.text)
        if 3 < len(text) < 40 and text == text.upper() and any(c.isalpha() for c in text):
            return text.title()
    return ""


def folien_titel(folie) -> str:
    """Titel = Text mit der groessten Schriftgroesse, bei Gleichstand der oberste."""
    kandidaten = []
    for shape in folie.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        groesse = 0.0
        for absatz in shape.text_frame.paragraphs:
            for lauf in absatz.runs:
                if lauf.font.size:
                    groesse = max(groesse, lauf.font.size.pt)
        kandidaten.append((groesse, shape.top or 0, saeubern(text.split("\n")[0])))
    if not kandidaten:
        return ""
    kandidaten.sort(key=lambda k: (-k[0], k[1]))
    return kandidaten[0][2][:70]


def teil_nummer(folie):
    """Nummer des Abschnitts, falls die Folie ein Teiler ist (Text 'TEIL 2')."""
    for shape in folie.shapes:
        if not shape.has_text_frame:
            continue
        treffer = re.fullmatch(r"TEIL\s*(\d+)", saeubern(shape.text_frame.text).upper())
        if treffer:
            return int(treffer.group(1))
    return None


def metadaten_lesen():
    from pptx import Presentation

    prs = Presentation(PPTX)
    roh = []
    for nummer, folie in enumerate(prs.slides, start=1):
        titel = folien_titel(folie)
        teil = teil_nummer(folie)
        notiz = ""
        if folie.has_notes_slide:
            notiz = folie.notes_slide.notes_text_frame.text.strip()
        if re.fullmatch(r"Notebook \d+", titel):
            titel = "Hands-on \u00b7 " + titel
        if teil:
            titel = f"Teil {teil} \u00b7 {titel}"
        roh.append({"n": nummer, "t": titel or f"Folie {nummer}", "k": notiz,
                    "teil": teil, "kick": kicker(folie)})

    # Gleichlautende Titel mit dem Kicker unterscheiden (z. B. Aha-Moment vs. Fazit)
    gesehen = {}
    for eintrag in roh:
        schluessel = eintrag["t"].rstrip(".")
        if schluessel in gesehen and eintrag["kick"]:
            eintrag["t"] = f"{eintrag['kick']}: {schluessel}"
        gesehen[schluessel] = True

    kapitel = "Einstieg"
    daten = []
    for eintrag in roh:
        if eintrag["teil"]:
            kapitel = eintrag["t"]
        daten.append({"n": eintrag["n"], "t": eintrag["t"], "k": eintrag["k"], "kap": kapitel})

    ohne = [d["n"] for d in daten if not d["k"]]
    if ohne:
        print(f"  Hinweis: keine Sprechernotiz auf Folie {ohne}")
    return daten


def pruefe_werkzeuge():
    fehlt = [w for w in ("soffice", "pdftoppm") if not shutil.which(w)]
    if fehlt:
        sys.exit(f"Fehlt im PATH: {', '.join(fehlt)}. Siehe Hinweis oben im Skript.")


def bilder_erzeugen():
    from PIL import Image

    TMP.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(TMP), str(PPTX)],
                   check=True, capture_output=True)
    pdf = TMP / (PPTX.stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-r", "96", "-scale-to-x", str(BREITE), "-scale-to-y", "-1",
                    str(pdf), str(TMP / "f")], check=True)

    OUT_IMG.mkdir(parents=True, exist_ok=True)
    for alt in OUT_IMG.glob("*.jpg"):
        alt.unlink()

    gesamt = 0
    seiten = sorted(TMP.glob("f-*.png"))
    for seite in seiten:
        nummer = int(re.search(r"f-(\d+)", seite.name).group(1))
        ziel = OUT_IMG / f"{nummer:02d}.jpg"
        Image.open(seite).convert("RGB").save(
            ziel, "JPEG", quality=QUALITAET, optimize=True, progressive=True)
        gesamt += ziel.stat().st_size
    shutil.rmtree(TMP, ignore_errors=True)
    print(f"  {len(seiten)} Bilder erzeugt ({gesamt / 1024 / 1024:.1f} MB) -> {OUT_IMG.relative_to(ROOT)}")
    return len(seiten)


def presenter_aktualisieren(daten):
    html = PRESENTER.read_text(encoding="utf-8")
    neu = json.dumps(daten, ensure_ascii=False, separators=(",", ":"))
    html, anzahl = re.subn(r"const S = \[.*?\];", f"const S = {neu};", html, count=1, flags=re.S)
    if anzahl != 1:
        sys.exit("Datenblock in praesentation.html nicht gefunden - bitte manuell pruefen.")
    PRESENTER.write_text(html, encoding="utf-8")
    print(f"  Titel und Notizen fuer {len(daten)} Folien eingesetzt")


if __name__ == "__main__":
    if not PPTX.exists():
        sys.exit(f"Nicht gefunden: {PPTX}")
    nur_notizen = "--nur-notizen" in sys.argv

    print("Praesentation wird aufgebaut ...")
    anzahl_bilder = None
    if not nur_notizen:
        pruefe_werkzeuge()
        anzahl_bilder = bilder_erzeugen()

    daten = metadaten_lesen()
    if anzahl_bilder is not None and anzahl_bilder != len(daten):
        print(f"  Achtung: {anzahl_bilder} Bilder, aber {len(daten)} Folien in der PPTX.")
    presenter_aktualisieren(daten)

    print("\nGliederung:")
    letztes = None
    for d in daten:
        if d["kap"] != letztes:
            print(f"  -- {d['kap']}")
            letztes = d["kap"]
        print(f"     {d['n']:2d}  {d['t']}")
    print("\nFertig. docs/praesentation.html ist auf dem aktuellen Stand.")
